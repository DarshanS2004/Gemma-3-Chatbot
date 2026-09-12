from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document

DATA_DIR = Path("data")
ASSET_DIR = DATA_DIR / "generated_assets"
IMAGE_DIR = ASSET_DIR / "images"
FILE_DIR = ASSET_DIR / "files"


def ensure_asset_dirs() -> None:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    FILE_DIR.mkdir(parents=True, exist_ok=True)


def slugify(value: str, fallback: str = "asset") -> str:
    compact = re.sub(r"[^a-zA-Z0-9]+", "-", value.strip().lower()).strip("-")
    return compact[:48] or fallback


def strip_code_fences(text: str) -> str:
    cleaned = (text or "").strip()
    if cleaned.startswith("```") and cleaned.endswith("```"):
        lines = cleaned.splitlines()
        if len(lines) >= 2:
            return "\n".join(lines[1:-1]).strip()
    if cleaned.startswith("~~~") and cleaned.endswith("~~~"):
        lines = cleaned.splitlines()
        if len(lines) >= 2:
            return "\n".join(lines[1:-1]).strip()
    return cleaned


def detect_image_prompt(prompt: str, mode: str | None = None) -> bool:
    lowered = (prompt or "").strip().lower()
    if mode == "imagegen":
        return True
    verbs = r"(generate|create|make|design|draw|render)"
    nouns = r"(image|photo|picture|illustration|art|poster|portrait|wallpaper|logo)"
    return bool(re.search(rf"\b{verbs}\b.*\b{nouns}\b", lowered))


def infer_file_format(prompt: str) -> str | None:
    lowered = (prompt or "").strip().lower()
    if not lowered:
        return None

    direct_map = {
        "docx": "docx",
        "word": "docx",
        "pdf": "pdf",
        "csv": "csv",
        "json": "json",
        "txt": "txt",
        "text file": "txt",
    }
    for needle, file_format in direct_map.items():
        if needle in lowered:
            return file_format

    if "downloadable format" in lowered or "downloadable file" in lowered or "generate a file" in lowered:
        if any(token in lowered for token in ("dataset", "table", "rows", "records", "reviews", "sentiment")):
            return "csv"
        if any(token in lowered for token in ("report", "proposal", "document", "assignment")):
            return "docx"
        return "pdf"

    return None


def detect_file_request(prompt: str) -> bool:
    lowered = (prompt or "").strip().lower()
    if not lowered:
        return False
    verb_hit = any(token in lowered for token in ("generate", "create", "make", "build", "prepare", "export", "give me"))
    noun_hit = any(token in lowered for token in ("file", "pdf", "csv", "docx", "word", "json", "txt", "downloadable"))
    return verb_hit and noun_hit


def build_file_generation_prompt(user_prompt: str, file_format: str) -> str:
    templates = {
        "csv": (
            "Create a realistic CSV file that satisfies the request below. "
            "Return only raw CSV text with a header row. No markdown fences, no explanation.\n\n"
            f"Request: {user_prompt}"
        ),
        "json": (
            "Create a valid JSON document that satisfies the request below. "
            "Return only raw JSON. No markdown fences, no explanation.\n\n"
            f"Request: {user_prompt}"
        ),
        "txt": (
            "Create a clean plain-text file that satisfies the request below. "
            "Use headings and lists when helpful, but return only the document text.\n\n"
            f"Request: {user_prompt}"
        ),
        "docx": (
            "Create well-structured document content for a downloadable DOCX that satisfies the request below. "
            "Use Markdown-style headings and bullet lists where helpful. Return only the document content.\n\n"
            f"Request: {user_prompt}"
        ),
        "pdf": (
            "Create well-structured document content for a downloadable PDF that satisfies the request below. "
            "Use Markdown-style headings and bullet lists where helpful. Return only the document content.\n\n"
            f"Request: {user_prompt}"
        ),
    }
    return templates[file_format]


def make_asset_filename(prompt: str, extension: str) -> str:
    base = slugify(prompt, fallback="asset")
    return f"{base}-{uuid4().hex[:10]}.{extension}"


def save_generated_image(image_bytes: bytes, prompt: str, image_format: str = "png") -> dict[str, str]:
    ensure_asset_dirs()
    extension = "jpg" if image_format == "jpeg" else image_format
    filename = make_asset_filename(prompt, extension)
    path = IMAGE_DIR / filename
    path.write_bytes(image_bytes)
    return {
        "filename": filename,
        "url": f"/api/assets/images/{filename}",
        "download_url": f"/api/assets/images/{filename}?download=1",
        "mime_type": f"image/{image_format}",
    }


def _normalize_csv_text(text: str) -> str:
    cleaned = strip_code_fences(text)
    lines = [line.rstrip() for line in cleaned.splitlines() if line.strip()]
    return "\n".join(lines).strip()


def _normalize_json_text(text: str) -> str:
    cleaned = strip_code_fences(text)
    try:
        parsed = json.loads(cleaned)
        return json.dumps(parsed, indent=2, ensure_ascii=True)
    except json.JSONDecodeError:
        return cleaned


def _normalize_document_text(text: str) -> str:
    cleaned = strip_code_fences(text)
    return cleaned.strip()


def _docx_bytes_from_text(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            document.add_paragraph("")
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(stripped)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _pdf_bytes_from_text(text: str) -> bytes:
    lines = []
    for raw_line in text.splitlines():
        stripped = raw_line.rstrip()
        if not stripped:
            lines.append("")
            continue
        wrapped = re.findall(r".{1,90}(?:\s+|$)", stripped) or [stripped]
        lines.extend(chunk.strip() for chunk in wrapped if chunk.strip())

    pages: list[list[str]] = []
    page: list[str] = []
    for line in lines:
        page.append(line)
        if len(page) >= 42:
            pages.append(page)
            page = []
    if page:
        pages.append(page)
    if not pages:
        pages = [["Generated file"]]

    objects: list[bytes] = []

    def add_object(data: str | bytes) -> int:
        payload = data.encode("latin-1") if isinstance(data, str) else data
        objects.append(payload)
        return len(objects)

    font_id = add_object("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids: list[int] = []
    content_ids: list[int] = []

    for page_lines in pages:
        content_stream_lines = ["BT", "/F1 11 Tf", "50 790 Td", "16 TL"]
        first = True
        for line in page_lines:
            escaped = _pdf_escape(line or " ")
            if first:
                content_stream_lines.append(f"({escaped}) Tj")
                first = False
            else:
                content_stream_lines.append(f"T* ({escaped}) Tj")
        content_stream_lines.append("ET")
        stream_body = "\n".join(content_stream_lines).encode("latin-1")
        content_id = add_object(
            b"<< /Length " + str(len(stream_body)).encode("ascii") + b" >>\nstream\n" + stream_body + b"\nendstream"
        )
        content_ids.append(content_id)

    pages_root_id = len(objects) + 1
    for content_id in content_ids:
        page_obj = (
            f"<< /Type /Page /Parent {pages_root_id} 0 R /MediaBox [0 0 612 842] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        )
        page_ids.append(add_object(page_obj))

    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    add_object(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>")
    catalog_id = add_object(f"<< /Type /Catalog /Pages {pages_root_id} 0 R >>")

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, payload in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("latin-1"))
        pdf.extend(payload)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("latin-1"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    pdf.extend(
        (
            f"trailer << /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF"
        ).encode("latin-1")
    )
    return bytes(pdf)


def save_generated_file(prompt: str, file_format: str, content: str) -> dict[str, str]:
    ensure_asset_dirs()
    filename = make_asset_filename(prompt, file_format)
    path = FILE_DIR / filename

    if file_format == "csv":
        path.write_text(_normalize_csv_text(content), encoding="utf-8")
        mime_type = "text/csv"
    elif file_format == "json":
        path.write_text(_normalize_json_text(content), encoding="utf-8")
        mime_type = "application/json"
    elif file_format == "txt":
        path.write_text(_normalize_document_text(content), encoding="utf-8")
        mime_type = "text/plain"
    elif file_format == "docx":
        path.write_bytes(_docx_bytes_from_text(_normalize_document_text(content)))
        mime_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif file_format == "pdf":
        path.write_bytes(_pdf_bytes_from_text(_normalize_document_text(content)))
        mime_type = "application/pdf"
    else:
        raise ValueError(f"Unsupported file format: {file_format}")

    return {
        "filename": filename,
        "url": f"/api/assets/files/{filename}",
        "download_url": f"/api/assets/files/{filename}?download=1",
        "mime_type": mime_type,
        "format": file_format.upper(),
    }


def asset_path(asset_kind: str, filename: str) -> Path:
    ensure_asset_dirs()
    base_dir = IMAGE_DIR if asset_kind == "images" else FILE_DIR
    path = (base_dir / Path(filename).name).resolve()
    if path.parent != base_dir.resolve() or not path.exists():
        raise FileNotFoundError(filename)
    return path
